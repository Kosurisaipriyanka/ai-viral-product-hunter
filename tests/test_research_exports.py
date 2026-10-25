from io import BytesIO

from openpyxl import load_workbook

from research_exports import create_excel_report, create_pdf_report


def sample_run():
    return {
        "id": 3,
        "run_type": "automated_discovery",
        "title": "Pet gadgets",
        "created_at": "2026-10-25T10:00:00+00:00",
        "payload": {
            "niche": "pet gadgets",
            "items": [
                {"rank": 1, "product_name": "Interactive cat toy", "opportunity_score": 62},
                {"rank": 2, "product_name": "Pet grooming glove", "opportunity_score": 48},
            ],
            "note": "Curated ideas only",
        },
    }


def test_excel_report_has_summary_and_results_sheets():
    content = create_excel_report(sample_run())
    workbook = load_workbook(BytesIO(content), read_only=True)
    assert workbook.sheetnames == ["Summary", "Results"]
    rows = list(workbook["Results"].values)
    assert rows[1][1] == "Interactive cat toy"


def test_pdf_report_is_generated():
    content = create_pdf_report(sample_run())
    assert content.startswith(b"%PDF-")
    assert len(content) > 500
